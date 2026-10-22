"""
generate.py - builds the synthetic Excel files in this folder (forecast, inventory, product master, summary).

Every printer model is given one stock SCENARIO, so the allocation rules can be tried on every case:
  SHORTAGE_SEVERE  inventory is about 30-50% of the total forecast   -> CONSTRAINED (weighted 25% baseline)
  SHORTAGE_MILD    inventory is about 80-95% of the total forecast   -> CONSTRAINED
  EXACT            inventory equals the total forecast               -> SUFFICIENT (every region gets its demand)
  SURPLUS_MILD     inventory is about 110-130% of the total forecast -> SUFFICIENT, some stock left unallocated
  SURPLUS_LARGE    inventory is about 200-300% of the total forecast -> SUFFICIENT, a lot left unallocated
  STOCK_OUT        no inventory at all                               -> CONSTRAINED, everyone gets 0
  NO_DEMAND        forecast is 0 in every region, stock exists      -> SUFFICIENT, all stock unallocated
  PDF_EXAMPLE      the one-pager's "Printer Model A": 100 units, demand 30/30/30/30 -> 25/25/25/25 baseline

The file layouts are exactly the upload templates the application validates
(backend/app/src/main/java/com/pifa/ingestion/ExcelFiles.java):
  Forecast  : sheet "Forecast"  - Product_name, Forecast_Month, North, East, West, South, Comments
  Inventory : sheet "Inventory" - Product_name, Available_Qty, As_Of_Date, Planning_Month
  Products  : sheet "Products"  - Product_name, Aliases, Active (Admin > Master data > Import)

Run (needs openpyxl):  python generate.py
Output is deterministic (fixed random seed), so re-running gives identical numbers.
"""
import datetime as dt
import random

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

FORECAST_MONTH = "2027-01"                      # the cycle the files are for
AS_OF = dt.date(2026, 10, 21)                   # inventory snapshot date (must not be in the future)
REGIONS = ["North", "East", "West", "South"]
random.seed(20270101)

# (model, family, typical monthly demand per region, scenario)
MODELS = [
    ("LaserJet Pro M110", "Mono laser", 120, "SHORTAGE_SEVERE"),
    ("LaserJet Pro M210", "Mono laser", 95, "SHORTAGE_MILD"),
    ("LaserJet Pro M310", "Mono laser", 70, "EXACT"),
    ("LaserJet Pro M410", "Mono laser", 45, "SURPLUS_MILD"),
    ("LaserJet Enterprise M610", "Mono laser", 25, "SURPLUS_LARGE"),
    ("Color LaserJet C220", "Colour laser", 80, "SHORTAGE_MILD"),
    ("Color LaserJet C320", "Colour laser", 60, "SHORTAGE_SEVERE"),
    ("Color LaserJet C420", "Colour laser", 40, "EXACT"),
    ("Color LaserJet C520", "Colour laser", 30, "SURPLUS_MILD"),
    ("Color LaserJet Enterprise C720", "Colour laser", 12, "STOCK_OUT"),
    ("InkTank 310", "Ink tank", 150, "SHORTAGE_MILD"),
    ("InkTank 315", "Ink tank", 140, "EXACT"),
    ("InkTank 410 Wireless", "Ink tank", 110, "SURPLUS_MILD"),
    ("InkTank 415 Wireless", "Ink tank", 90, "SHORTAGE_SEVERE"),
    ("InkTank 515 All-in-One", "Ink tank", 65, "SURPLUS_LARGE"),
    ("SmartTank 580", "Ink tank", 55, "EXACT"),
    ("DeskJet 2330", "Inkjet", 160, "SHORTAGE_MILD"),
    ("DeskJet 2720", "Inkjet", 130, "SURPLUS_MILD"),
    ("DeskJet Plus 4120", "Inkjet", 85, "SHORTAGE_SEVERE"),
    ("OfficeJet Pro 8020", "Inkjet", 50, "EXACT"),
    ("OfficeJet Pro 9010", "Inkjet", 35, "SURPLUS_LARGE"),
    ("MFP Laser M140w", "Laser MFP", 75, "SHORTAGE_MILD"),
    ("MFP Laser M236sdw", "Laser MFP", 55, "EXACT"),
    ("MFP Laser M438n", "Laser MFP", 30, "SURPLUS_MILD"),
    ("MFP Color M283fdw", "Laser MFP", 40, "SHORTAGE_SEVERE"),
    ("PageWide Pro 477dw", "PageWide", 15, "STOCK_OUT"),
    ("DesignJet T230 Plotter", "Large format", 6, "SURPLUS_LARGE"),
    ("DesignJet T650 Plotter", "Large format", 4, "NO_DEMAND"),
    ("ScanJet Pro 2000 s2", "Scanner", 20, "SURPLUS_MILD"),
    ("Printer Model A", "Example (one-pager)", 30, "PDF_EXAMPLE"),
]

# Inventory as a share of total forecast for each scenario (low, high).
RATIO = {
    "SHORTAGE_SEVERE": (0.30, 0.50),
    "SHORTAGE_MILD": (0.80, 0.95),
    "SURPLUS_MILD": (1.10, 1.30),
    "SURPLUS_LARGE": (2.00, 3.00),
}

# Regional demand shape: North and West sell a bit more, South a bit less (varied per model below).
REGION_FACTOR = {"North": 1.15, "East": 1.00, "West": 1.05, "South": 0.80}

COMMENTS = {
    "SHORTAGE_SEVERE": "Supply constrained - vendor allocation cut",
    "STOCK_OUT": "Awaiting shipment",
    "NO_DEMAND": "End of life - no new orders",
    "PDF_EXAMPLE": "One-pager example: 100 units, equal demand",
}


def demand_for(base, scenario, region):
    """Forecast for one model and region: base x regional factor x +/-15% noise (0 for NO_DEMAND)."""
    if scenario == "NO_DEMAND":
        return 0
    if scenario == "PDF_EXAMPLE":
        return 30
    noise = random.uniform(0.85, 1.15)
    return max(1, round(base * REGION_FACTOR[region] * noise))


def inventory_for(total, scenario, base):
    """Available stock for one model given its total forecast and scenario."""
    if scenario == "EXACT":
        return total
    if scenario == "STOCK_OUT":
        return 0
    if scenario == "NO_DEMAND":
        return base * 10                       # stock exists although nobody forecasts demand
    if scenario == "PDF_EXAMPLE":
        return 100
    lo, hi = RATIO[scenario]
    return round(total * random.uniform(lo, hi))


def header(ws, names):
    """Writes a bold header row and sensible column widths."""
    for i, n in enumerate(names, start=1):
        c = ws.cell(row=1, column=i, value=n)
        c.font = Font(bold=True)
        ws.column_dimensions[get_column_letter(i)].width = 34 if i == 1 else 16
    ws.freeze_panes = "A2"


rows = []
for model, family, base, scenario in MODELS:
    by_region = {r: demand_for(base, scenario, r) for r in REGIONS}
    total = sum(by_region.values())
    stock = inventory_for(total, scenario, base)
    rows.append((model, family, scenario, by_region, total, stock))

# ---- Forecast file
wb = Workbook()
ws = wb.active
ws.title = "Forecast"
header(ws, ["Product_name", "Forecast_Month"] + REGIONS + ["Comments"])
for model, _, scenario, by_region, _, _ in rows:
    ws.append([model, FORECAST_MONTH] + [by_region[r] for r in REGIONS] + [COMMENTS.get(scenario)])
wb.save(f"FORECAST_{FORECAST_MONTH}.xlsx")

# ---- Inventory file
wb = Workbook()
ws = wb.active
ws.title = "Inventory"
header(ws, ["Product_name", "Available_Qty", "As_Of_Date", "Planning_Month"])
for model, _, _, _, _, stock in rows:
    ws.append([model, stock, AS_OF, FORECAST_MONTH])
    ws.cell(row=ws.max_row, column=3).number_format = "yyyy-mm-dd"
wb.save(f"INVENTORY_{AS_OF.isoformat()}.xlsx")

# ---- Product master (import before uploading the two files above)
wb = Workbook()
ws = wb.active
ws.title = "Products"
header(ws, ["Product_name", "Aliases", "Active"])
for model, *_ in rows:
    ws.append([model, None, "yes"])
wb.save("PRODUCTS_master.xlsx")

# ---- Summary for people (not for upload): what each model should do in the allocation
EXPECT = {
    "SHORTAGE_SEVERE": "CONSTRAINED - split by 25% baseline, regions short",
    "SHORTAGE_MILD": "CONSTRAINED - split by 25% baseline, small shortfall",
    "EXACT": "SUFFICIENT - every region gets exactly its demand",
    "SURPLUS_MILD": "SUFFICIENT - full demand, some stock unallocated",
    "SURPLUS_LARGE": "SUFFICIENT - full demand, much stock unallocated",
    "STOCK_OUT": "CONSTRAINED - nothing to allocate (all 0)",
    "NO_DEMAND": "SUFFICIENT - nothing needed, all stock unallocated",
    "PDF_EXAMPLE": "CONSTRAINED - 25/25/25/25 baseline (one-pager example)",
}
FILL = {"SHORT": "F9D6D3", "EXACT": "D9F2E0", "SURPLUS": "DBEAFE"}
wb = Workbook()
ws = wb.active
ws.title = "Summary"
header(ws, ["Product_name", "Family", "Scenario"] + REGIONS + ["Total forecast", "Available inventory", "Gap (inv - fcst)", "Coverage %", "Expected allocation"])
ws.column_dimensions["L"].width = 52
for model, family, scenario, by_region, total, stock in rows:
    cover = round(stock / total * 100) if total else None
    ws.append([model, family, scenario] + [by_region[r] for r in REGIONS] + [total, stock, stock - total, cover, EXPECT[scenario]])
    kind = "SHORT" if stock < total else "EXACT" if stock == total else "SURPLUS"
    for col in range(1, 13):
        ws.cell(row=ws.max_row, column=col).fill = PatternFill("solid", fgColor=FILL[kind])
# Scenario counts at the bottom.
ws.append([])
ws.append(["Scenario counts"])
ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
for s in EXPECT:
    ws.append([s, sum(1 for r in rows if r[2] == s)])
wb.save("SCENARIO_SUMMARY.xlsx")

short = sum(1 for r in rows if r[5] < r[4])
exact = sum(1 for r in rows if r[5] == r[4])
print(f"{len(rows)} models: {short} inventory < forecast, {exact} equal, {len(rows) - short - exact} inventory > forecast")
